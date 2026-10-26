"""Build excel/option_pricer.xlsx: a Black-Scholes pricer in native Excel formulas.

No macros are needed. Run from the project root:  python excel/build_workbook.py
"""

from pathlib import Path
import sys

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent / "option_pricer.xlsx"

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from option_pricer import bs_price, crr_price, lsm_price, mc_price  # noqa: E402

INPUT_FILL = PatternFill("solid", fgColor="FFF4CC")
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="BBBBBB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def header(ws, cell, text):
    ws[cell] = text
    ws[cell].fill = HEADER_FILL
    ws[cell].font = HEADER_FONT


def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "Pricer"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 46

    ws["A1"] = "Black-Scholes option pricer"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Edit the yellow cells. Everything else is computed with Excel formulas."
    ws["A2"].font = Font(italic=True, color="666666")

    # --- Inputs (named ranges make the formulas readable) -------------------
    header(ws, "A4", "Inputs")
    inputs = [
        ("Spot (S)", 100, "Spot", "Underlying price"),
        ("Strike (K)", 100, "Strike", ""),
        ("Maturity in years (T)", 1, "Maturity", "e.g. 0.5 = 6 months"),
        ("Risk-free rate (r)", 0.05, "Rate", "Continuously compounded"),
        ("Volatility (sigma)", 0.20, "Vol", "Annualised"),
        ("Dividend yield (q)", 0.0, "DivYield", "Continuous"),
        ("Type (Call / Put)", "Call", "OptType", "Drop-down list"),
    ]
    from openpyxl.workbook.defined_name import DefinedName
    for i, (label, value, name, note) in enumerate(inputs):
        row = 5 + i
        ws.cell(row=row, column=1, value=label)
        c = ws.cell(row=row, column=2, value=value)
        c.fill, c.border = INPUT_FILL, BOX
        ws.cell(row=row, column=3, value=note).font = Font(color="888888")
        ref = f"Pricer!$B${row}"
        dn = DefinedName(name, attr_text=ref)
        try:
            wb.defined_names[name] = dn          # openpyxl >= 3.1
        except TypeError:
            wb.defined_names.append(dn)          # older openpyxl
    for r in (8, 9, 10):
        ws.cell(row=r, column=2).number_format = "0.00%"
    dv = DataValidation(type="list", formula1='"Call,Put"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B11")

    # --- Intermediate terms -------------------------------------------------
    header(ws, "A13", "Intermediate")
    ws["A14"], ws["B14"] = "d1", "=(LN(Spot/Strike)+(Rate-DivYield+Vol^2/2)*Maturity)/(Vol*SQRT(Maturity))"
    ws["A15"], ws["B15"] = "d2", "=B14-Vol*SQRT(Maturity)"
    ws["A16"], ws["B16"] = "Sign (+1 call, -1 put)", '=IF(OptType="Call",1,-1)'
    for r in (14, 15):
        ws.cell(row=r, column=2).number_format = "0.0000"

    # --- Outputs -------------------------------------------------------------
    header(ws, "A18", "Results")
    w = "B16"  # +1 / -1
    outputs = [
        ("Price", f"={w}*(Spot*EXP(-DivYield*Maturity)*NORMSDIST({w}*B14)"
                  f"-Strike*EXP(-Rate*Maturity)*NORMSDIST({w}*B15))", ""),
        ("Delta", f"={w}*EXP(-DivYield*Maturity)*NORMSDIST({w}*B14)", "per 1 unit of spot"),
        ("Gamma", "=EXP(-DivYield*Maturity)*EXP(-B14^2/2)/SQRT(2*PI())/(Spot*Vol*SQRT(Maturity))", ""),
        ("Vega (per 1 vol point)", "=Spot*EXP(-DivYield*Maturity)*EXP(-B14^2/2)/SQRT(2*PI())*SQRT(Maturity)/100",
         "change for +1% volatility"),
        ("Theta (per calendar day)",
         f"=(-Spot*EXP(-DivYield*Maturity)*EXP(-B14^2/2)/SQRT(2*PI())*Vol/(2*SQRT(Maturity))"
         f"-{w}*Rate*Strike*EXP(-Rate*Maturity)*NORMSDIST({w}*B15)"
         f"+{w}*DivYield*Spot*EXP(-DivYield*Maturity)*NORMSDIST({w}*B14))/365",
         "annual theta / 365"),
        ("Rho (per 1% rate)", f"={w}*Strike*Maturity*EXP(-Rate*Maturity)*NORMSDIST({w}*B15)/100",
         "change for +1% rate"),
        ("Forward value S.e^(-qT) - K.e^(-rT)",
         "=Spot*EXP(-DivYield*Maturity)-Strike*EXP(-Rate*Maturity)", "C - P should equal this"),
    ]
    for i, (label, formula, note) in enumerate(outputs):
        row = 19 + i
        ws.cell(row=row, column=1, value=label).font = BOLD if i == 0 else Font()
        c = ws.cell(row=row, column=2, value=formula)
        c.number_format, c.border = "0.0000", BOX
        ws.cell(row=row, column=3, value=note).font = Font(color="888888")
    ws["B19"].font = BOLD

    # --- Price vs spot table and chart --------------------------------------
    ch = wb.create_sheet("Chart data")
    ch.column_dimensions["A"].width = 12
    ch.column_dimensions["B"].width = 14
    ch.column_dimensions["C"].width = 14
    for col, text in zip("ABC", ("Spot", "Option price", "Payoff")):
        header(ch, f"{col}1", text)
    n = 41
    for i in range(n):
        row = 2 + i
        # Spot from 50% to 150% of strike
        ch[f"A{row}"] = f"=Strike*(0.5+{i}/{n - 1})"
        s = f"A{row}"
        d1 = f"((LN({s}/Strike)+(Rate-DivYield+Vol^2/2)*Maturity)/(Vol*SQRT(Maturity)))"
        d2 = f"({d1}-Vol*SQRT(Maturity))"
        ch[f"B{row}"] = (f"=Pricer!$B$16*({s}*EXP(-DivYield*Maturity)*NORMSDIST(Pricer!$B$16*{d1})"
                         f"-Strike*EXP(-Rate*Maturity)*NORMSDIST(Pricer!$B$16*{d2}))")
        ch[f"C{row}"] = f"=MAX(Pricer!$B$16*({s}-Strike),0)"
        for col in "ABC":
            ch[f"{col}{row}"].number_format = "0.00"

    chart = LineChart()
    chart.title = "Option price vs spot (today) and payoff at maturity"
    chart.x_axis.title = "Spot"
    chart.y_axis.title = "Value"
    chart.height, chart.width = 9, 17
    chart.add_data(Reference(ch, min_col=2, min_row=1, max_col=3, max_row=n + 1), titles_from_data=True)
    chart.set_categories(Reference(ch, min_col=1, min_row=2, max_row=n + 1))
    chart.x_axis.number_format = "0"
    chart.x_axis.tickLblSkip = 5
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    ws.add_chart(chart, "E4")

    monte_carlo_sheet(wb)

    # --- Notes ----------------------------------------------------------------
    notes = wb.create_sheet("Notes")
    notes.column_dimensions["A"].width = 100
    for i, line in enumerate([
        "Model: Black-Scholes-Merton for European options, continuous dividend yield.",
        "Uses only long-standing functions (NORMSDIST, LN, EXP), so it works in any Excel version and LibreOffice.",
        "Greeks: vega and rho per 1 percentage point, theta per calendar day (annual / 365).",
        "The 'Monte Carlo' sheet holds values precomputed in Python (python excel/build_workbook.py);"
        " they do not follow the yellow input cells.",
        "American options are priced in Python (CRR tree, Longstaff-Schwartz); results are on the Monte Carlo sheet.",
        "Optional: vba/BlackScholes.bas adds a =BS_PRICE(...) user function (see README).",
    ], start=1):
        notes[f"A{i}"] = line
        notes[f"A{i}"].alignment = Alignment(wrap_text=True)

    wb.save(OUT)
    print(f"Workbook written to {OUT}")


def monte_carlo_sheet(wb):
    """Static sheet: Monte Carlo results computed in Python, plus a native line chart."""
    S, K, T, r, sigma = 100.0, 100.0, 1.0, 0.05, 0.20
    bs = bs_price(S, K, T, r, sigma, "call")
    mc = wb.create_sheet("Monte Carlo", 1)
    mc["A1"] = "Monte Carlo convergence, ATM European call (precomputed in Python)"
    mc["A1"].font = Font(bold=True, size=14)
    mc["A2"] = (f"S={S:.0f}, K={K:.0f}, T={T:.0f}y, r={r:.0%}, vol={sigma:.0%}, seed=N (one independent run per row). "
                "Static values: re-run python excel/build_workbook.py to refresh.")
    mc["A2"].font = Font(italic=True, color="666666")

    cols = ["Paths (N)", "Black-Scholes", "Plain MC", "Plain low 95%", "Plain high 95%",
            "Antithetic + CV", "Reduced low 95%", "Reduced high 95%", "Std error ratio"]
    widths = [12, 14, 12, 14, 14, 16, 16, 16, 15]
    for j, (text, w) in enumerate(zip(cols, widths), start=1):
        cell = mc.cell(row=4, column=j, value=text)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, horizontal="center")
        mc.column_dimensions[cell.column_letter].width = w
    ns = [500, 1_000, 2_000, 5_000, 10_000, 20_000, 50_000, 100_000, 200_000, 500_000]
    for i, n in enumerate(ns):
        row = 5 + i
        plain = mc_price(S, K, T, r, sigma, "call", n_paths=n, antithetic=False,
                         control_variate=False, seed=n)
        red = mc_price(S, K, T, r, sigma, "call", n_paths=n, seed=n)
        values = [n, bs, plain.price, plain.ci_low, plain.ci_high, red.price, red.ci_low,
                  red.ci_high, plain.std_error / red.std_error]
        for j, v in enumerate(values, start=1):
            c = mc.cell(row=row, column=j, value=round(v, 6) if j > 1 else v)
            c.border = BOX
            c.number_format = "#,##0" if j == 1 else ("0.0" if j == 9 else "0.0000")
    last = 4 + len(ns)

    chart = LineChart()
    chart.title = "MC price and 95% interval vs number of paths"
    chart.x_axis.title = "Number of paths"
    chart.y_axis.title = "Call price"
    chart.height, chart.width = 9, 20
    chart.add_data(Reference(mc, min_col=2, min_row=4, max_col=8, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(mc, min_col=1, min_row=5, max_row=last))
    colors = ["000000", "9AA5B1", "C9CFD6", "C9CFD6", "1F6FEB", "8DB6F5", "8DB6F5"]
    dashes = ["dash", None, "sysDot", "sysDot", None, "sysDot", "sysDot"]
    for series, color, dash in zip(chart.series, colors, dashes):
        series.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.width = 22000
        if dash:
            series.graphicalProperties.line.dashStyle = dash
        series.smooth = False
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.y_axis.scaling.min = 9.5
    chart.y_axis.scaling.max = 11.5
    mc.add_chart(chart, f"A{last + 3}")

    # American put: tree vs Longstaff-Schwartz
    top = last + 23
    mc[f"A{top}"] = "American put, Longstaff & Schwartz (2001) case: S=36, K=40, r=6%, vol=20%, T=1"
    mc[f"A{top}"].font = BOLD
    rows = [("European put (Black-Scholes)", bs_price(36, 40, 1, 0.06, 0.20, "put"), None),
            ("American put, CRR tree 2,000 steps", crr_price(36, 40, 1, 0.06, 0.20, "put",
                                                              steps=2000, american=True), None)]
    lsm = lsm_price(36, 40, 1, 0.06, 0.20, "put", n_paths=100_000, steps=50, seed=1)
    rows.append(("American put, LSM 100,000 paths, 50 dates", lsm.price, lsm.std_error))
    rows.append(("Paper, finite differences (reference grid)", 4.478, None))
    rows.append(("Paper, LSM simulation", 4.472, 0.010))
    for j, text in enumerate(["Method", "Price", "Std error"], start=1):
        cell = mc.cell(row=top + 1, column=j, value=text)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for i, (label, price, se) in enumerate(rows):
        row = top + 2 + i
        mc.cell(row=row, column=1, value=label)
        c = mc.cell(row=row, column=2, value=round(price, 6))
        c.number_format, c.border = "0.0000", BOX
        if se is not None:
            c = mc.cell(row=row, column=3, value=round(se, 6))
            c.number_format, c.border = "0.0000", BOX
    mc.cell(row=top + 7, column=1, value=(
        "The tree is stable at 4.4867 from 1,000 to 20,000 steps; the paper's 4.478 comes from a coarser "
        "grid. LSM is noisy (s.e. above) and slightly biased low (estimated exercise rule).")).font = Font(
        italic=True, color="666666")
    mc.column_dimensions["A"].width = 42
    mc.freeze_panes = "A5"


if __name__ == "__main__":
    build()
